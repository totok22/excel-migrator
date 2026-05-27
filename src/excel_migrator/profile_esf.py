"""FSEC ESF specific migration rules."""

from __future__ import annotations

import copy
from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from .core import CellAction, copyable_value, hyperlink_target_from_value, is_input_cell, norm_text


def _existing(ws: Any, row: int, col: int) -> Any:
    if row < 1 or col < 1:
        return None
    return ws._cells.get((row, col))


def _can_write(cell: Any, overwrite: bool) -> bool:
    from .core import is_placeholder, is_instructional
    if cell is None or isinstance(cell, MergedCell):
        return False
    if isinstance(cell.value, str) and cell.value.startswith("="):
        return False
    if overwrite:
        return True
    return cell.value is None or is_placeholder(cell.value) or is_instructional(cell.value)


def _is_warning_value(value: Any) -> bool:
    """Detect formula-generated warning/status text that should not be migrated."""
    from .core import is_status_value
    return is_status_value(value)


def _copy_if_possible(source_ws: Any, target_ws: Any, src_coord: str, tgt_coord: str,
                     actions: list[CellAction], method: str, overwrite: bool) -> None:
    from openpyxl.utils import coordinate_to_tuple
    s = source_ws._cells.get(coordinate_to_tuple(src_coord))
    tr, tc = coordinate_to_tuple(tgt_coord)
    t = target_ws._cells.get((tr, tc))
    if isinstance(t, MergedCell):
        for merged in target_ws.merged_cells.ranges:
            if tgt_coord in merged:
                t = target_ws._cells.get((merged.min_row, merged.min_col))
                break
    if copyable_value(s) and _can_write(t, overwrite):
        # Skip formula-generated status text like "Warning", "Incomplete"
        if _is_warning_value(s.value):
            return
        t.value = s.value
        if s.hyperlink:
            t._hyperlink = copy.copy(s.hyperlink)
            t._hyperlink.ref = t.coordinate
        else:
            t.hyperlink = hyperlink_target_from_value(s.value)
        actions.append(CellAction(target_ws.title, s.coordinate, t.coordinate, method, s.value))


def grounding_target_rows(ws: Any, label_col: int) -> list[int]:
    rows: list[int] = []
    for row in range(4, ws.max_row + 1):
        cell = _existing(ws, row, label_col)
        if cell is not None and not isinstance(cell, MergedCell) and (cell.value is not None or is_input_cell(cell)):
            if row == 4 or row % 10 == 4:
                rows.append(row)
    return rows


def migrate_grounding(source_wb: Any, target_wb: Any, overwrite: bool) -> list[CellAction]:
    name = "接地 Grounding"
    if name not in source_wb.sheetnames or name not in target_wb.sheetnames:
        return []
    source_ws = source_wb[name]
    target_ws = target_wb[name]
    actions: list[CellAction] = []

    target_by_part: dict[str, int] = {}
    for row in range(4, target_ws.max_row + 1):
        cell = _existing(target_ws, row, 1)
        part = norm_text(cell.value if cell is not None else None)
        if part:
            target_by_part[part] = row

    for row in range(4, source_ws.max_row + 1):
        cell = _existing(source_ws, row, 1)
        part = norm_text(cell.value if cell is not None else None)
        if part and part in target_by_part:
            tr = target_by_part[part]
            for s_col, t_col in ((2, 2), (3, 3), (4, 4)):
                _copy_if_possible(
                    source_ws, target_ws,
                    f"{get_column_letter(s_col)}{row}",
                    f"{get_column_letter(t_col)}{tr}",
                    actions, "grounding-part", overwrite,
                )

    source_rows = [
        r for r in range(4, source_ws.max_row + 1)
        if copyable_value(_existing(source_ws, r, 7)) or copyable_value(_existing(source_ws, r, 8))
    ]
    target_rows = grounding_target_rows(target_ws, 7)
    for sr, tr in zip(source_rows, target_rows):
        for s_col, t_col in ((7, 7), (8, 8)):
            _copy_if_possible(
                source_ws, target_ws,
                f"{get_column_letter(s_col)}{sr}",
                f"{get_column_letter(t_col)}{tr}",
                actions, "grounding-carbon-row", overwrite,
            )

    source_rows = [
        r for r in range(4, source_ws.max_row + 1)
        if any(copyable_value(_existing(source_ws, r, c)) for c in (11, 12, 13, 14))
    ]
    target_rows = grounding_target_rows(target_ws, 11)
    for sr, tr in zip(source_rows, target_rows):
        for s_col, t_col in ((11, 11), (12, 12), (13, 13), (14, 14)):
            _copy_if_possible(
                source_ws, target_ws,
                f"{get_column_letter(s_col)}{sr}",
                f"{get_column_letter(t_col)}{tr}",
                actions, "grounding-enclosure-row", overwrite,
            )

    return actions


def grounding_image_targeter(source_ws: Any, target_ws: Any, source_img: Any) -> Any | str | None:
    """Targets used by image migration for ESF sheets."""
    from .images import image_anchor_start

    row, col = image_anchor_start(source_img)

    if source_ws.title == "备用电池箱 Spare Accumulator":
        return _spare_accumulator_image_anchor(source_img, row, col)

    if source_ws.title != "接地 Grounding":
        return None

    if col == 3:
        cell = _existing(source_ws, row, 1)
        part = norm_text(cell.value if cell is not None else None)
        for tr in range(4, target_ws.max_row + 1):
            tcell = _existing(target_ws, tr, 1)
            if norm_text(tcell.value if tcell is not None else None) == part:
                return f"C{tr}"
    if col == 8:
        used = [
            r for r in range(4, source_ws.max_row + 1)
            if copyable_value(_existing(source_ws, r, 7)) or copyable_value(_existing(source_ws, r, 8))
        ]
        target_rows = grounding_target_rows(target_ws, 7)
        if row in used:
            idx = used.index(row)
            if idx < len(target_rows):
                return f"H{target_rows[idx]}"
    if col == 13:
        used = [
            r for r in range(4, source_ws.max_row + 1)
            if any(copyable_value(_existing(source_ws, r, c)) for c in (11, 12, 13, 14))
        ]
        target_rows = grounding_target_rows(target_ws, 11)
        if row in used:
            idx = used.index(row)
            if idx < len(target_rows):
                return f"M{target_rows[idx]}"
    return None


def _shift_image_anchor_rows(source_img: Any, delta: int, *, min_start_row: int | None = None,
                             max_end_row: int | None = None) -> Any:
    anchor = copy.deepcopy(source_img.anchor)
    start = getattr(anchor, "_from", None)
    if start is not None:
        start.row += delta
        if min_start_row is not None:
            start.row = max(start.row, min_start_row - 1)

    end = getattr(anchor, "to", None)
    if end is not None:
        end.row += delta
        if max_end_row is not None:
            end.row = min(end.row, max_end_row - 1)
    return anchor


def _spare_accumulator_image_anchor(source_img: Any, row: int, col: int) -> Any | None:
    """Move old spare-accumulator picture anchors onto the shifted 2026 template blocks."""
    # The 2026 Spare Accumulator template moved the picture blocks down while the
    # old workbooks keep images anchored to the earlier rows. Preserve each
    # image's relative columns/offsets so side-by-side pictures stay side-by-side.
    if 11 <= row <= 30 and 39 <= col <= 44:
        return _shift_image_anchor_rows(source_img, 12, min_start_row=23, max_end_row=43)
    if 34 <= row <= 53 and 1 <= col <= 53:
        return _shift_image_anchor_rows(source_img, 12, min_start_row=47, max_end_row=65)
    if 66 <= row <= 84 and 1 <= col <= 17:
        return _shift_image_anchor_rows(source_img, 12, min_start_row=78, max_end_row=96)
    return None


def fix_known_template_formula_breaks(wb: Any) -> list[str]:
    fixes: list[str] = []
    if "总览 Overview" in wb.sheetnames:
        ws = wb["总览 Overview"]
        if ws["K19"].value == "=#REF!":
            ws["K19"] = "='电池箱 Accumulator'!U20"
            fixes.append("总览 Overview!K19: =#REF! -> ='电池箱 Accumulator'!U20")
    if "其他 Others" in wb.sheetnames:
        ws = wb["其他 Others"]
        if isinstance(ws["Z32"].value, str) and "#REF!" in ws["Z32"].value:
            ws["Z32"] = '=IF(ISBLANK(Y32),IncompleteText,IF(Y32="非锂电池",OKText,AI28))'
            fixes.append('其他 Others!Z32: removed #REF! and linked lithium LVS check to AI28')
    return fixes


ESF_SKIP_SHEETS = {"接地 Grounding"}

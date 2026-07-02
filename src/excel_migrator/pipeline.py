"""Top-level migration pipeline shared by CLI and web server."""

from __future__ import annotations

import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from .core import (
    CellAction,
    SheetCache,
    SkippedCell,
    clear_target_placeholders,
    migrate_generic_cells,
    set_recalc_on_open,
)
from .images import ImageAction, SkippedImage, migrate_images_generic
from .profile_esf import (
    ESF_SKIP_SHEETS,
    fix_known_template_formula_breaks,
    grounding_image_targeter,
    migrate_datasheet_links,
    migrate_grounding,
)
from .report import ReportInputs, write_excel_report
from .strict import build_strict_output, extra_empty_rows_not_in_template, restore_empty_cells_from_reference


ProgressFn = Callable[[str, int, int], None]


@dataclass
class MigrationOptions:
    source: Path
    template: Path
    output: Path
    excel_report: Path
    profile: str = "generic"  # "generic" | "esf"
    overwrite: bool = False
    include_images: bool = True
    keep_template_images: bool = False
    # Advanced settings
    context_threshold: float = 0.55
    fuzzy_threshold: float = 0.70
    image_margin: float = 0.92
    cross_sheet: bool = False
    filter_status: bool = True
    keep_instructional: bool = True


@dataclass
class MigrationResult:
    output: Path
    excel_report: Path
    cell_actions: list[CellAction] = field(default_factory=list)
    skipped_cells: list[SkippedCell] = field(default_factory=list)
    image_actions: list[ImageAction] = field(default_factory=list)
    skipped_images: list[SkippedImage] = field(default_factory=list)
    formula_fixes: list[str] = field(default_factory=list)
    strict_patched_cells: int = 0
    remaining_extra_empty_rows: int = 0
    elapsed_seconds: float = 0.0


def _build_caches(wb: Any) -> dict[str, SheetCache]:
    return {ws.title: SheetCache(ws) for ws in wb.worksheets}


def _normalize_output_for_excel(path: Path, template_path: Path) -> None:
    """Rewrite the workbook package once so Office Excel accepts the output."""
    wb = load_workbook(path)
    normalized: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmpf:
            normalized = Path(tmpf.name)
        wb.save(normalized)
    finally:
        wb.close()
    restore_empty_cells_from_reference(path, normalized, layout_path=template_path)
    try:
        normalized.replace(path)
    finally:
        if normalized is not None and normalized.exists():
            normalized.unlink()


def run(opts: MigrationOptions, progress: ProgressFn | None = None) -> MigrationResult:
    started = time.perf_counter()

    def emit(stage: str, current: int, total: int) -> None:
        if progress:
            progress(stage, current, total)

    emit("加载旧版", 1, 8)
    source_wb = load_workbook(opts.source)
    emit("加载新版模板", 2, 8)
    target_wb = load_workbook(opts.template)
    from .core import set_migration_config
    set_migration_config(
        context_threshold=opts.context_threshold,
        fuzzy_threshold=opts.fuzzy_threshold,
        image_margin=opts.image_margin,
        cross_sheet=opts.cross_sheet,
        filter_status=opts.filter_status,
        keep_instructional=opts.keep_instructional,
    )
    clear_target_placeholders(target_wb)

    cell_actions: list[CellAction] = []
    skipped_cells: list[SkippedCell] = []
    skip_sheets: set[str] = set()
    formula_fixes: list[str] = []

    if opts.profile == "esf":
        emit("处理 FSEC ESF 接地表", 3, 8)
        cell_actions.extend(migrate_grounding(source_wb, target_wb, opts.overwrite))
        skip_sheets |= ESF_SKIP_SHEETS
        formula_fixes.extend(fix_known_template_formula_breaks(target_wb))

    emit("预计算 worksheet 缓存", 4, 8)
    source_caches = _build_caches(source_wb)
    target_caches = _build_caches(target_wb)

    def cell_progress(stage: str, current: int, total: int) -> None:
        emit(f"通用迁移：{stage}", 5, 8)

    actions, skips = migrate_generic_cells(
        source_caches,
        target_caches,
        opts.overwrite,
        skip_sheets=skip_sheets,
        progress=cell_progress,
    )
    cell_actions.extend(actions)
    skipped_cells.extend(skips)

    if opts.profile == "esf":
        link_actions = migrate_datasheet_links(source_wb, target_wb, opts.overwrite)
        if link_actions:
            cell_actions.extend(link_actions)
            filled_links = {(a.sheet, a.target) for a in link_actions}
            skipped_cells = [s for s in skipped_cells if (s.sheet, s.coord) not in filled_links]

    image_actions: list[ImageAction] = []
    skipped_images: list[SkippedImage] = []
    if opts.include_images:
        emit("迁移图片", 6, 8)
        targeter = grounding_image_targeter if opts.profile == "esf" else None
        image_actions, skipped_images = migrate_images_generic(
            source_wb,
            target_wb,
            keep_template_images=opts.keep_template_images,
            custom_targeter=targeter,
        )

    set_recalc_on_open(target_wb)

    emit("写出新版输出", 7, 8)
    opts.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmpf:
        staging = Path(tmpf.name)
    try:
        target_wb.save(staging)
        strict_patched, name_fixes = build_strict_output(opts.template, staging, opts.output)
        formula_fixes.extend(name_fixes)
    finally:
        if staging.exists():
            staging.unlink()

    remaining_extra = sum(
        len(rows) for rows in extra_empty_rows_not_in_template(opts.template, opts.output).values()
    )

    elapsed = time.perf_counter() - started

    emit("生成报告", 8, 8)
    info = ReportInputs(
        source_path=opts.source,
        template_path=opts.template,
        output_path=opts.output,
        profile=opts.profile,
        overwrite=opts.overwrite,
        elapsed_seconds=elapsed,
        cell_actions=cell_actions,
        skipped_cells=skipped_cells,
        image_actions=image_actions,
        skipped_images=skipped_images,
        formula_fixes=formula_fixes,
        strict_patched_cells=strict_patched,
        remaining_extra_empty_rows=remaining_extra,
    )
    write_excel_report(opts.excel_report, info)

    return MigrationResult(
        output=opts.output,
        excel_report=opts.excel_report,
        cell_actions=cell_actions,
        skipped_cells=skipped_cells,
        image_actions=image_actions,
        skipped_images=skipped_images,
        formula_fixes=formula_fixes,
        strict_patched_cells=strict_patched,
        remaining_extra_empty_rows=remaining_extra,
        elapsed_seconds=elapsed,
    )

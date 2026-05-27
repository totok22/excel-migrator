"""Image migration."""

from __future__ import annotations

import copy
import io
from dataclasses import dataclass
from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import coordinate_to_tuple, get_column_letter


@dataclass
class ImageAction:
    sheet: str
    source_anchor: str
    target_anchor: str
    method: str
    size: str


@dataclass
class SkippedImage:
    sheet: str
    anchor: str
    reason: str


def image_anchor_start(img: Any) -> tuple[int, int]:
    anchor = img.anchor
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return 1, 1
    return marker.row + 1, marker.col + 1


def anchor_label(img: Any) -> str:
    row, col = image_anchor_start(img)
    return f"{get_column_letter(col)}{row}"


def anchor_start_label(anchor: Any) -> str:
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return "?"
    return f"{get_column_letter(marker.col + 1)}{marker.row + 1}"


def image_bytes(img: Any) -> bytes:
    if hasattr(img, "_data"):
        return img._data()
    ref = getattr(img, "ref", None)
    if hasattr(ref, "getvalue"):
        return ref.getvalue()
    if hasattr(ref, "read"):
        pos = ref.tell()
        ref.seek(0)
        data = ref.read()
        ref.seek(pos)
        return data
    raise ValueError("Cannot read image bytes")


def column_width_px(width: float | None) -> float:
    width = 8.43 if width is None else width
    return max(1.0, width * 7.0 + 5.0)


def row_height_px(height: float | None) -> float:
    height = 15.0 if height is None else height
    return max(1.0, height * 96.0 / 72.0)


def range_pixels(ws: Any, min_row: int, min_col: int, max_row: int, max_col: int) -> tuple[float, float]:
    width = 0.0
    for c in range(min_col, max_col + 1):
        d = ws.column_dimensions.get(get_column_letter(c))
        width += column_width_px(d.width if d is not None else None)
    height = 0.0
    for r in range(min_row, max_row + 1):
        d = ws.row_dimensions.get(r)
        height += row_height_px(d.height if d is not None else None)
    return width, height


def containing_range(ws: Any, coord: str) -> tuple[int, int, int, int]:
    for merged in ws.merged_cells.ranges:
        if coord in merged:
            return merged.min_row, merged.min_col, merged.max_row, merged.max_col
    row, col = coordinate_to_tuple(coord)
    return row, col, row, col


def fit_image_to_cell(img: Any, ws: Any, coord: str) -> None:
    from .core import _config
    min_row, min_col, max_row, max_col = containing_range(ws, coord)
    bw, bh = range_pixels(ws, min_row, min_col, max_row, max_col)
    margin = _config.image_margin
    bw *= margin
    bh *= margin
    if img.width <= 0 or img.height <= 0:
        return
    scale = min(bw / img.width, bh / img.height)
    if scale > 0:
        img.width = max(1, int(img.width * scale))
        img.height = max(1, int(img.height * scale))


def _compute_cell_offset_emu(ws: Any, row: int, col: int, img_width: int, img_height: int) -> tuple[int, int]:
    """Compute EMU offsets to center an image within its containing merged cell.

    Returns (col_offset_emu, row_offset_emu) to center the image.
    """
    from openpyxl.utils import get_column_letter as gcl
    coord = f"{gcl(col)}{row}"
    min_row, min_col, max_row, max_col = containing_range(ws, coord)
    cell_w, cell_h = range_pixels(ws, min_row, min_col, max_row, max_col)

    # Convert image dimensions from pixels to approximate size
    # Center horizontally and vertically within the cell
    dx = max(0, (cell_w - img_width) / 2)
    dy = max(0, (cell_h - img_height) / 2)

    # Convert pixels to EMU (1 pixel = 9525 EMU)
    EMU_PER_PX = 9525
    return int(dx * EMU_PER_PX), int(dy * EMU_PER_PX)


def add_image_copy(target_ws: Any, source_img: Any, anchor: Any, fit_coord: str | None = None) -> None:
    new_img = XLImage(io.BytesIO(image_bytes(source_img)))
    new_img.width = source_img.width
    new_img.height = source_img.height

    if fit_coord:
        # Scale image to fit within the target cell
        fit_image_to_cell(new_img, target_ws, fit_coord)
        target_ws.add_image(new_img, fit_coord)
    else:
        # Preserve original anchor exactly (keeps TwoCellAnchor type and all offsets)
        new_img.anchor = copy.deepcopy(anchor)
        target_ws.add_image(new_img)


def migrate_images_generic(
    source_wb: Any,
    target_wb: Any,
    keep_template_images: bool,
    custom_targeter=None,
) -> tuple[list[ImageAction], list[SkippedImage]]:
    """Copy images from old workbook to new workbook by anchor.

    custom_targeter(source_ws, target_ws, img) -> str | anchor | None:
    special target coord or copied anchor (used by FSEC ESF profile).
    """
    actions: list[ImageAction] = []
    skipped: list[SkippedImage] = []
    if not keep_template_images:
        for ws in target_wb.worksheets:
            ws._images = []
    for source_ws in source_wb.worksheets:
        if source_ws.title not in target_wb.sheetnames:
            continue
        target_ws = target_wb[source_ws.title]
        for img in getattr(source_ws, "_images", []):
            sa = anchor_label(img)
            try:
                ta = None
                if custom_targeter is not None:
                    ta = custom_targeter(source_ws, target_ws, img)
                if isinstance(ta, str) and ta:
                    add_image_copy(target_ws, img, img.anchor, fit_coord=ta)
                    actions.append(ImageAction(source_ws.title, sa, ta, "fit-cell", f"{img.width}x{img.height}"))
                elif ta:
                    add_image_copy(target_ws, img, ta)
                    actions.append(
                        ImageAction(source_ws.title, sa, anchor_start_label(ta), "custom-anchor", f"{img.width}x{img.height}")
                    )
                else:
                    add_image_copy(target_ws, img, img.anchor)
                    actions.append(ImageAction(source_ws.title, sa, sa, "same-anchor", f"{img.width}x{img.height}"))
            except Exception as exc:
                skipped.append(SkippedImage(source_ws.title, sa, str(exc)))
    return actions, skipped
